#!/usr/bin/env python3
"""Build an author-facing review sheet from manual_mapping_review.csv.

Reads validation/output/manual_mapping_review.csv (not modified) and writes
  validation/output/manual_mapping_review_author.xlsx
  validation/output/manual_mapping_review_author.csv
The author-assessment columns are left empty. Automatic flags are labelled as
automatic pre-analysis, never as human judgements.

Usage: python validation/build_author_review_sheet.py
"""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "validation" / "output" / "manual_mapping_review.csv"
OUT_XLSX = ROOT / "validation" / "output" / "manual_mapping_review_author.xlsx"
OUT_CSV = ROOT / "validation" / "output" / "manual_mapping_review_author.csv"

OPTIONS = [
    "1 - Mevcut metadata eşlemeyi destekliyor",
    "2 - Eşleme kurala uygun, ancak gerçek DLP anlamı için bağlam yetersiz",
    "3 - Eşleme hatalı",
    "4 - Karar verilemedi",
]

CATEGORY = {"A_EMAIL": "E-posta", "B_HTTP_CLOUD": "HTTP / bulut", "C_REMOVABLE_DEVICE": "Çıkarılabilir cihaz",
            "D_FILE": "Dosya"}
SUBGROUP = {
    "A": "Dış alıcı + ek (e-posta koruma kapsamı)",
    "B-random": "Bulut listesi eşleşmesi (rastgele)",
    "C": "Cihaz bağlantısı (USB koruma kapsamı)",
    "D1": "Path var, tanınan uzantı",
    "D2": "Path var, tanınmayan/eksik uzantı (OTHER)",
    "D3": "Path var, hassasiyet anahtar kelimesi eşleşti",
    "D4": "Yalnızca dosya adı, tanınan uzantı",
    "D5": "Yalnızca dosya adı, .doc uzantısı (OTHER)",
}

FLAG_TR = {
    "SENDER_NOT_INTERNAL_DOMAIN": "Gönderen adresi dtaa.com dışında. Kayıt, kurumsal e-posta ağ geçidi yerine kişisel "
                                  "e-posta trafiği olabilir; adaptör göndereni dikkate almıyor.",
    "MIXED_INTERNAL_EXTERNAL_RECIPIENTS": "Alıcılar arasında hem iç (dtaa.com) hem dış alan adı var; \"en az bir dış "
                                          "alıcı\" kuralı nedeniyle dış gönderim sayıldı.",
    "EXTERNAL_ONLY_VIA_BCC": "Dış alıcı yalnızca Bcc alanında.",
    "ATTACHMENT_CONTENT_NOT_IN_SOURCE": "Ek sayısı biliniyor; ekin adı, türü ve içeriği kaynakta yok. 'Restricted' "
                                        "sınıflandırması doğrulanamaz.",
    "NO_RECIPIENT_DOMAIN": "Alıcı alan adı çıkarılamadı.",
    "CLOUD_MATCH_IS_ACCESS_NOT_UPLOAD": "Host sabit bulut listesinde, ancak kayıt yalnızca URL erişimi. HTTP metodu, "
                                        "aktarılan bayt veya yükleme göstergesi yok; CloudUpload bir vekil (proxy) "
                                        "eşlemedir, dosya yüklemesi doğrulanmış değildir.",
    "TRANSFER_LIKE_HOST_NOT_ON_LIST": "Host adı dosya aktarımı çağrıştırıyor ama sabit bulut listesinde değil.",
    "PRIVATE_IP_HOST_MARKED_UNTRUSTED": "Host özel (iç) IP adresi, ancak tüm HTTP kayıtları gibi 'Untrusted' atandı.",
    "HOST_NOT_PARSEABLE": "URL'den host çıkarılamadı.",
    "SIZE_MISSING_DEFAULTS_TO_0": "Kaynakta boyut alanı yok; size_mb=0 atandı.",
    "DEVICE_TRUST_ASSUMED_UNMANAGED": "Kaynakta cihaz yönetim durumu yok; 'Unmanaged' her cihaza varsayım olarak "
                                      "atanır.",
    "NO_DEVICE_IDENTIFIER_IN_SOURCE": "Cihaz kimliği, üretici, seri numarası veya cihaz türü bilgisi yok.",
    "UNKNOWN_DEVICE_ACTION": "Cihaz eylemi Connect/Disconnect dışında.",
    "FILENAME_ONLY_SENSITIVITY_NOT_EVALUABLE": "Path yok, yalnızca dosya adı var; yol tabanlı hassasiyet kuralı "
                                               "uygulanamadı ve varsayılan FileOperation/Internal atandı.",
    "EXTENSION_OUTSIDE_TYPE_DOMAIN": "Uzantı tanımlı dosya türleri (PDF, DOCX, XLSX, CSV, ZIP, TXT) arasında değil; "
                                     "file_type=OTHER.",
    "NO_EXTENSION": "Dosyanın uzantısı yok; file_type=OTHER.",
    "SENSITIVE_BY_KEYWORD": "Yolda hassasiyet anahtar kelimesi eşleşti; SensitiveFile/Confidential yalnızca bu kelime "
                            "eşleşmesine dayanıyor, dosya içeriği doğrulanmadı.",
}

VERIFIABLE_TR = {
    "A_EMAIL": ("Gönderim eylemi; alıcı alan adları; dış alıcı varlığı (dtaa.com tek iç alan adı varsayımıyla); ek "
                "sayısı.",
                "Ek içeriği ve hassasiyeti; ek adı/türü; dtaa.com'un tek iç alan adı olduğu; gönderen cihaz."),
    "B_HTTP_CLOUD": ("URL erişimi; host adı; hostun sabit bulut listesinde olması.",
                     "Dosya yüklemesi; veri hassasiyeti; hedef güven düzeyi; boyut."),
    "C_REMOVABLE_DEVICE": ("Bağlanma/ayrılma eylemi; zaman damgası.",
                           "Cihazın yönetim durumu (Unmanaged varsayım); cihaz türü; cihaza veri yazılıp yazılmadığı."),
    "D_FILE": ("Dosya eylemi; dosya adı veya yolu; uzantı.",
               "Dosya içeriğinin hassasiyeti; gerçek dosya biçimi (yalnızca uzantı var); Path yoksa dizin konumu."),
}


def flags_text(codes):
    return "\n".join(f"• {FLAG_TR.get(code, code)}" for code in codes)


def source_metadata(r):
    g = r["group"]
    if g == "A_EMAIL":
        return (f"{r['source_hostname_or_recipient']}\nHam alıcılar: {r['source_recipients_raw']}\n"
                f"Ek sayısı (Attachments): {r['source_attachments']}")
    if g == "B_HTTP_CLOUD":
        return f"Host: {r['source_hostname_or_recipient']}\nURL: {r['source_url']}"
    if g == "C_REMOVABLE_DEVICE":
        return r["source_device_metadata"]
    return (f"Filename: {r['source_filename'] or '(yok)'}\nPath: {r['source_path'] or '(yok)'}")


def build_rows(df):
    rows = []
    for _, r in df.iterrows():
        codes = [c for c in r["flags"].split(";") if c]
        amb_codes = [c for c in r["ambiguity_reasons"].split(";") if c]
        other = [c for c in codes if c not in amb_codes]
        subgroup = SUBGROUP.get(r["subgroup"], "")
        if r["subgroup"].startswith("B-min1/"):
            subgroup = f"Bulut listesi eşleşmesi (liste girdisi başına 1: {r['subgroup'].split('/', 1)[1]})"
        ver, nover = VERIFIABLE_TR[r["group"]]
        rows.append({
            "İnceleme ID": r["review_id"],
            "Kategori": CATEGORY[r["group"]],
            "Alt grup": subgroup,
            "SPEDIA kayıt ID": r["source_record_id"],
            "Zaman damgası": r["source_timestamp"],
            "Kaynak eylem (Action)": r["source_action"],
            "Kaynak metadata": source_metadata(r),
            "Kaynak boyut bilgisi": r["source_size_metadata"],
            "Adaptör eşlemesi": (f"channel={r['adapter_channel']}; content_type={r['adapter_content_type']}; "
                                 f"data_classification={r['adapter_data_classification']}; "
                                 f"destination_trust={r['adapter_destination_trust']}; "
                                 f"device_trust={r['adapter_device_trust']}; file_type={r['adapter_file_type']}; "
                                 f"size_mb={r['adapter_size_mb']}"),
            "Koruma kapsamı": r["protection_scope"],
            "Uygulanan adaptör kuralı": r["mapping_rule"],
            "Otomatik kontrol: kurala uygunluk": ("Kurala uygun" if r["rule_consistent"] == "yes"
                                                  else f"Uyumsuz ({r['rule_consistency_mismatch_fields']})"),
            "Otomatik belirsizlik bayrağı": "Evet" if r["ambiguous"] == "yes" else "Hayır",
            "Otomatik belirsizlik gerekçesi": flags_text(amb_codes) if amb_codes else "—",
            "Diğer otomatik notlar": flags_text(other) if other else "—",
            "Kaynakta doğrulanabilen": ver,
            "Kaynakta doğrulanamayan": nover,
            "Yazar değerlendirmesi": "",
            "Yazar açıklaması": "",
        })
    return rows


WIDTHS = [10, 14, 28, 22, 14, 14, 48, 22, 46, 24, 46, 16, 12, 48, 48, 34, 38, 34, 40]


def main():
    df = pd.read_csv(SOURCE, encoding="utf-8-sig", dtype=str, keep_default_na=False)
    if len(df) != 100:
        raise SystemExit(f"Expected 100 records, found {len(df)}")
    if (df["human_review"] != "").any() or (df["reviewer_notes"] != "").any():
        raise SystemExit("Source already contains human review entries; refusing to build a blank sheet.")
    rows = build_rows(df)
    headers = list(rows[0].keys())
    pd.DataFrame(rows, columns=headers).to_csv(OUT_CSV, index=False, encoding="utf-8-sig")

    base = Font(name="Arial", size=9)
    bold = Font(name="Arial", size=9, bold=True)
    head_font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2F4F6F")
    auto_fill = PatternFill("solid", fgColor="EDEDED")
    input_fill = PatternFill("solid", fgColor="FFF2CC")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()

    # ---- Sheet 1: instructions
    info = wb.active
    info.title = "Açıklama"
    lines = [
        ("RuleScope-DLP — SPEDIA eşleme incelemesi (100 kayıt)", True),
        ("", False),
        ("Amaç: Adaptörün SPEDIA metadata'sından ürettiği DLP eşlemelerini yazarın kayıt kayıt incelemesi.", False),
        ("Kaynak: validation/output/manual_mapping_review.csv (değiştirilmedi). Seçim: seed 20261007, kategori "
         "başına 25 kayıt; yöntem için manual_mapping_review_method.md.", False),
        ("", False),
        ("Nasıl doldurulur", True),
        ("• Yalnızca 'İnceleme' sayfasındaki SARI sütunları doldurun: 'Yazar değerlendirmesi' (açılır liste) ve "
         "'Yazar açıklaması' (serbest metin).", False),
        ("• GRİ sütunlar otomatik ön analizdir; insan kararı DEĞİLDİR. Değerlendirmenizi bunlardan bağımsız verin.",
         False),
        ("• 2, 3 veya 4 seçildiğinde açıklama alanına kısa gerekçe yazın.", False),
        ("", False),
        ("Yazar değerlendirmesi seçenekleri", True),
        *[(f"   {option}", False) for option in OPTIONS],
        ("", False),
        ("Otomatik sütunların anlamı", True),
        ("• Otomatik kontrol: kurala uygunluk — bağımsız bir denetleyicinin, kaynak metadata'dan belgelenmiş "
         "adaptör kuralıyla aynı değeri türetip türetmediği. Kuralın DLP açısından doğru olduğunu göstermez.", False),
        ("• Otomatik belirsizlik bayrağı — betiğin kayda özgü bir durum nedeniyle (ör. gönderen alan adı, eksik "
         "Path) eşlemenin gerçek anlamını belirsiz saydığı kayıtlar. Bir sezgisel kuraldır, yargı değildir.", False),
        ("• Diğer otomatik notlar — belirsizlik sayılmayan ama kaynakta doğrulanamayan özellikler (ör. bulut "
         "eşleşmesinin yükleme kanıtı olmaması, cihaz güveninin varsayım olması).", False),
        ("", False),
        ("Örnek (yalnızca biçim gösterimi; veri satırı değildir)", True),
        ("   Yazar değerlendirmesi: 2 - Eşleme kurala uygun, ancak gerçek DLP anlamı için bağlam yetersiz", False),
        ("   Yazar açıklaması: Dış alıcı ve ek kaynakta görülüyor; gönderen kişisel adres, kurumsal kanal "
         "belirsiz.", False),
    ]
    for i, (text, is_bold) in enumerate(lines, start=1):
        cell = info.cell(row=i, column=1, value=text)
        cell.font = Font(name="Arial", size=12 if i == 1 else 10, bold=is_bold)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    info.column_dimensions["A"].width = 120

    # ---- Sheet 2: review table
    ws = wb.create_sheet("İnceleme")
    input_cols = {"Yazar değerlendirmesi", "Yazar açıklaması"}
    auto_cols = {"Otomatik kontrol: kurala uygunluk", "Otomatik belirsizlik bayrağı",
                 "Otomatik belirsizlik gerekçesi", "Diğer otomatik notlar"}
    for c, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=header)
        cell.font, cell.fill, cell.border = head_font, head_fill, border
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(c)].width = WIDTHS[c - 1]
    ws.cell(row=1, column=headers.index("Otomatik belirsizlik bayrağı") + 1).comment = Comment(
        "Betik tarafından hesaplanan otomatik ön analiz; insan değerlendirmesi değildir.", "build_author_review_sheet")
    for r, row in enumerate(rows, start=2):
        for c, header in enumerate(headers, start=1):
            cell = ws.cell(row=r, column=c, value=row[header])
            cell.font = bold if header == "İnceleme ID" else base
            cell.alignment, cell.border = wrap, border
            if header in input_cols:
                cell.fill = input_fill
            elif header in auto_cols:
                cell.fill = auto_fill
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    ws.row_dimensions[1].height = 42

    lists = wb.create_sheet("Listeler")
    for i, option in enumerate(OPTIONS, start=1):
        lists.cell(row=i, column=1, value=option).font = base
    lists.sheet_state = "hidden"
    eval_col = get_column_letter(headers.index("Yazar değerlendirmesi") + 1)
    dv = DataValidation(type="list", formula1="=Listeler!$A$1:$A$4", allow_blank=True,
                        error="Lütfen listedeki dört seçenekten birini seçin.", errorTitle="Geçersiz değer",
                        prompt="1–4 arasından seçin", promptTitle="Yazar değerlendirmesi")
    dv.showErrorMessage = True
    dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(f"{eval_col}2:{eval_col}{len(rows) + 1}")

    # ---- Sheet 3: summary (formulas)
    sm = wb.create_sheet("Özet")
    cat_col = get_column_letter(headers.index("Kategori") + 1)
    amb_col = get_column_letter(headers.index("Otomatik belirsizlik bayrağı") + 1)
    con_col = get_column_letter(headers.index("Otomatik kontrol: kurala uygunluk") + 1)
    last = len(rows) + 1
    rng = lambda col: f"'İnceleme'!${col}$2:${col}${last}"  # noqa: E731
    summary_headers = ["Kategori", "Kayıt", "Otomatik: kurala uygun", "Otomatik: belirsiz bayrağı",
                       "Yazar: 1 destekliyor", "Yazar: 2 bağlam yetersiz", "Yazar: 3 hatalı",
                       "Yazar: 4 karar verilemedi", "Yazar: değerlendirilmemiş"]
    for c, h in enumerate(summary_headers, start=1):
        cell = sm.cell(row=1, column=c, value=h)
        cell.font, cell.fill, cell.border = head_font, head_fill, border
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        sm.column_dimensions[get_column_letter(c)].width = 16 if c > 1 else 22
    cats = list(CATEGORY.values())
    for i, cat in enumerate(cats, start=2):
        sm.cell(row=i, column=1, value=cat)
        sm.cell(row=i, column=2, value=f'=COUNTIFS({rng(cat_col)},$A{i})')
        sm.cell(row=i, column=3, value=f'=COUNTIFS({rng(cat_col)},$A{i},{rng(con_col)},"Kurala uygun")')
        sm.cell(row=i, column=4, value=f'=COUNTIFS({rng(cat_col)},$A{i},{rng(amb_col)},"Evet")')
        for k, option in enumerate(OPTIONS):
            sm.cell(row=i, column=5 + k, value=f'=COUNTIFS({rng(cat_col)},$A{i},{rng(eval_col)},"{option}")')
        sm.cell(row=i, column=9, value=f"=B{i}-SUM(E{i}:H{i})")
    total_row = len(cats) + 2
    sm.cell(row=total_row, column=1, value="Toplam")
    for c in range(2, 10):
        col = get_column_letter(c)
        sm.cell(row=total_row, column=c, value=f"=SUM({col}2:{col}{total_row - 1})")
    for row in sm.iter_rows(min_row=2, max_row=total_row, max_col=9):
        for cell in row:
            cell.font = bold if cell.row == total_row else base
            cell.border = border
    note = sm.cell(row=total_row + 2, column=1,
                   value="C–D sütunları otomatik ön analizdir; E–I sütunları yalnızca 'İnceleme' sayfasındaki yazar "
                         "girişlerinden hesaplanır.")
    note.font = Font(name="Arial", size=9, italic=True)

    wb.move_sheet("Özet", offset=-1)
    wb.save(OUT_XLSX)
    print(f"Wrote {OUT_XLSX.name} and {OUT_CSV.name} ({len(rows)} records)")


if __name__ == "__main__":
    main()
